# DOCX mirror of hwpbuild.Builder, used for page-count estimation and as a fallback deliverable.
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from hwpbuild import CHARS, PARAS, F_MJ, F_HD, F_GO, BF_HEAD, BF_BOX, BF_SHADE

FONTMAP = {F_MJ: ("휴먼명조", "NanumMyeongjo"), F_HD: ("HY헤드라인M", "NanumSquare ExtraBold"),
           F_GO: ("맑은 고딕", "NanumGothic")}
ALIGN = {"JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY, "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
         "CENTER": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT}
FILL = {BF_HEAD: "DCE6F2", BF_BOX: "F2F2F2", BF_SHADE: "F2F2F2"}


class DocxBuilder:
    def __init__(self, proxy_fonts=True):
        self.proxy = proxy_fonts
        self.doc = Document()
        s = self.doc.sections[0]
        self._setup(s)
        self._pending_break = False
        st = self.doc.styles["Normal"]
        st.font.size = Pt(12)

    def _setup(self, s):
        s.page_width = Mm(210); s.page_height = Mm(297)
        s.left_margin = Mm(30); s.right_margin = Mm(30)
        s.top_margin = Mm(20); s.bottom_margin = Mm(15)
        s.header_distance = Mm(15); s.footer_distance = Mm(15)

    def _font(self, run, ck):
        font, size, bold, color = CHARS[ck]
        name = FONTMAP[font][1 if self.proxy else 0]
        run.font.name = name
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), name)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color[1:])

    def _fmt(self, p, style):
        align, left, ind, prev, nxt, ls, keep, tab = PARAS[style]
        pf = p.paragraph_format
        pf.alignment = ALIGN[align]
        pf.left_indent = Pt(left)
        pf.first_line_indent = Pt(ind)
        pf.space_before = Pt(prev); pf.space_after = Pt(nxt)
        pf.keep_with_next = bool(keep)
        pf.widow_control = False
        if tab == "toc":
            from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
            pf.tab_stops.add_tab_stop(Mm(148), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        return ls

    def para(self, runs, style="body", page_break=False, _container=None):
        if isinstance(runs, str):
            runs = [(runs, None)]
        cont = _container if _container is not None else self.doc
        if _container is None:
            self._fresh = False
        p = cont.add_paragraph()
        ls = self._fmt(p, style)
        if (page_break or self._pending_break) and _container is None:
            p.paragraph_format.page_break_before = True
            self._pending_break = False
        default_char = {"h1": "h1", "h1top": "h1", "h2": "h2", "h3": "h3", "h4": "h4", "cap": "cap",
                        "capfig": "cap", "note": "note", "tblc": "tbl", "tbll": "tbl", "tblj": "tbl",
                        "toc1": "toc1", "toc2": "toc2", "toc3": "toc3", "boxp": "boxb", "boxt": "boxt"}.get(style, "body")
        maxsize = 0
        for text, ck in runs or [("", None)]:
            ck = ck or default_char
            r = p.add_run(text)
            self._font(r, ck)
            maxsize = max(maxsize, CHARS[ck][1])
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(maxsize * ls / 100.0)
        return p

    def page_break(self):
        if not getattr(self, "_fresh", False):
            self._pending_break = True

    def _shade(self, cell, color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def table(self, rows, widths=None, header_rows=1, aligns=None, merges=(), font="tbl",
              row_shade=(), cell_bf=None, para_style=None, repeat=True):
        from hwpbuild import split_marks
        nrows = len(rows); ncols = max(len(r) for r in rows)
        widths = widths or [1] * ncols
        tot = sum(widths)
        if self._pending_break:
            self.para("", "tblc")
        self._fresh = False
        t = self.doc.add_table(rows=nrows, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._element.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        tw = tblPr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblPr.append(tw)
        tw.set(qn("w:w"), str(int(150 * 56.7))); tw.set(qn("w:type"), "dxa")
        grid = t._element.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(150 * 56.7 * w / tot)))
        for (r0, c0, rs, cs) in merges:
            a = t.cell(r0, c0); b = t.cell(r0 + rs - 1, c0 + cs - 1)
            a.merge(b)
        covered = set()
        for (r0, c0, rs, cs) in merges:
            for rr in range(r0, r0 + rs):
                for cc in range(c0, c0 + cs):
                    if (rr, cc) != (r0, c0):
                        covered.add((rr, cc))
        for r in range(nrows):
            for c in range(ncols):
                cell = t.cell(r, c)
                cell.width = Mm(150 * widths[c] / tot)
                if (r, c) in covered:
                    continue
                is_head = r < header_rows
                bf = BF_HEAD if is_head else (BF_SHADE if r in row_shade else None)
                if cell_bf and (r, c) in cell_bf:
                    bf = cell_bf[(r, c)]
                if bf in FILL:
                    self._shade(cell, FILL[bf])
                txt = rows[r][c] if c < len(rows[r]) else ""
                cell._element.remove(cell.paragraphs[0]._element)
                if isinstance(txt, list):
                    for pst2, runs2 in txt:
                        self.para(runs2, pst2, _container=cell)
                    continue
                al = "c" if is_head else (aligns[c] if aligns else "c")
                pst = para_style or {"c": "tblc", "l": "tbll", "j": "tblj"}[al]
                for line in str(txt).split("\n"):
                    ck = (font + "b") if is_head and font + "b" in CHARS else font
                    runs = []
                    for seg, red in split_marks(line):
                        k = ck
                        if red:
                            k = "tblred" if font == "tbl" else "notered"
                        runs.append((seg, k))
                    self.para(runs, pst, _container=cell)
        return t

    def box(self, paras):
        return self.table([[paras]], widths=[1], header_rows=0, cell_bf={(0, 0): BF_BOX}, repeat=False)

    def picture(self, path, width_mm, height_mm=None):
        p = self.doc.add_paragraph()
        self._fmt(p, "fig")
        if self._pending_break:
            p.paragraph_format.page_break_before = True; self._pending_break = False
        p.add_run().add_picture(path, width=Mm(width_mm))
        return p

    def new_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup(s)
        self._pending_break = False
        self._fresh = True
        return s

    def save(self, path):
        self.doc.save(path)
